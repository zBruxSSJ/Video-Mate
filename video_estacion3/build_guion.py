"""Genera Guion_Desafio_Parametros.docx a partir del guion de las 8 escenas."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PERSON_COLORS = {
    "Bruno Ferro": RGBColor(0xB4, 0x7B, 0x0E),
    "Fabrizio Osorio": RGBColor(0x0D, 0x94, 0x88),
    "Sebastián Cardoso": RGBColor(0xC2, 0x41, 0x0C),
    "Eugenia Pérez": RGBColor(0x6D, 0x28, 0xD9),
    "Todos / Bruno": RGBColor(0xB4, 0x7B, 0x0E),
}


def fmt_time(s):
    m = int(s) // 60
    sec = int(s) % 60
    return f"{m}:{sec:02d}"


SCENES = [
    dict(num=1, title="Intro", dur=14, speaker="Bruno Ferro",
         text=[
             "¡Hola a todos! Bienvenidos a nuestro video de Matemática, tercero BATIN, "
             "Estación 3: Problema 2.",
             "Tenemos que hallar a y b para que g de x sea continua en todos los reales.",
             "Somos Bruno Ferro, Fabrizio Osorio, Sebastián Cardoso y Eugenia Pérez, "
             "del San Ignacio de Loyola School. ¡Vamos a resolverlo!",
         ]),
    dict(num=2, title="Condiciones de continuidad", dur=20, speaker="Fabrizio Osorio",
         text=[
             "Antes de resolver el problema, repasemos las tres condiciones que debe "
             "cumplir una función para ser continua en un punto x igual a a.",
             "Primera condición: g de a debe existir. El valor de la función en x igual "
             "a a tiene que estar definido.",
             "Segunda condición: el límite cuando x tiende a a debe existir. Esto significa "
             "que el límite lateral izquierdo y el límite lateral derecho tienen que ser iguales.",
             "Tercera condición: ese límite tiene que ser igual a g de a.",
             "Si se cumplen las tres condiciones, la función es continua en ese punto.",
         ]),
    dict(num=3, title="Planteo del problema", dur=33, speaker="Fabrizio Osorio",
         text=[
             "Ahora sí, planteemos nuestro problema. Tenemos la función g de x, "
             "definida por partes:",
             "x más 1, si x es menor que 1;  a x más b, si x está entre 1 y 3;  "
             "y 3x, si x es mayor o igual que 3.",
             "Cada uno de estos tramos es una función lineal, así que cada tramo "
             "ya es continuo dentro de su propio intervalo.",
             "El problema está en los puntos donde la función cambia de fórmula: "
             "x igual a 1 y x igual a 3.",
             "Esos son los únicos puntos en riesgo de discontinuidad, y ahí es donde "
             "vamos a aplicar las tres condiciones que acabamos de repasar.",
         ]),
    dict(num=4, title="Condición en x = 1", dur=22, speaker="Sebastián Cardoso",
         text=[
             "Empecemos analizando la continuidad en x igual a 1.",
             "Primero, calculamos el límite lateral izquierdo usando el tramo x más 1, "
             "porque es el que se usa cuando x se acerca a 1 por la izquierda: "
             "el límite, cuando x tiende a 1 por la izquierda, de x más 1, es igual a 2.",
             "Segundo, calculamos el valor de la función en x igual a 1. Como x igual a 1 "
             "pertenece al tramo a x más b, tenemos que g de 1 es igual a a por 1 más b, "
             "es decir, a más b.",
             "Tercero, para que la función sea continua, el límite tiene que ser igual "
             "al valor de la función. Entonces igualamos: a más b, igual a 2. "
             "Esta es nuestra Ecuación 1.",
         ]),
    dict(num=5, title="Condición en x = 3", dur=20, speaker="Sebastián Cardoso",
         text=[
             "Hagamos exactamente lo mismo en el otro punto de unión: x igual a 3.",
             "Primero, el límite lateral izquierdo, usando el tramo a x más b: "
             "el límite, cuando x tiende a 3 por la izquierda, de a x más b, "
             "es igual a 3a más b.",
             "Segundo, el valor de la función en x igual a 3. Como x igual a 3 pertenece "
             "al tramo 3x, tenemos que g de 3 es igual a 3 por 3, que es 9.",
             "Tercero, igualamos el límite con el valor de la función: "
             "3a más b, igual a 9. Esta es nuestra Ecuación 2.",
         ]),
    dict(num=6, title="Resolver el sistema", dur=28, speaker="Eugenia Pérez",
         text=[
             "Ahora tenemos un sistema de dos ecuaciones con dos incógnitas: "
             "a más b igual a 2, y 3a más b igual a 9.",
             "Para resolverlo, restamos la Ecuación 2 menos la Ecuación 1: "
             "3a más b, menos a más b, igual a 9 menos 2.",
             "Al simplificar, los b se cancelan y nos queda: 2a igual a 7. "
             "Entonces, a es igual a siete medios, que es 3 coma 5.",
             "Ahora sustituimos ese valor en la Ecuación 1: siete medios más b, igual a 2. "
             "Despejando b, obtenemos: b igual a 2 menos siete medios, que es "
             "menos tres medios, o menos 1 coma 5.",
             "Entonces, nuestra solución es: a igual a siete medios, y b igual a "
             "menos tres medios.",
         ]),
    dict(num=7, title="Verificación gráfica", dur=31, speaker="Eugenia Pérez",
         text=[
             "Comprobemos gráficamente que todo esto tiene sentido. Vamos a dibujar "
             "los tres tramos de la función.",
             "Este es el primer tramo, x más 1, en color turquesa. Este es el segundo "
             "tramo, siete medios x menos tres medios, en naranja. Y este es el tercer "
             "tramo, 3x, en rojo.",
             "Miren cómo los tramos se unen exactamente en los puntos (1, 2) y (3, 9), "
             "sin ningún salto ni ningún agujero.",
             "Repasando el panel: g de a existe en ambos puntos, los límites laterales "
             "son iguales, y el límite coincide con g de a en los dos casos.",
             "Se cumplen las tres condiciones, así que podemos afirmar que g es "
             "continua en todo su dominio.",
         ]),
    dict(num=8, title="Cierre", dur=19, speaker="Todos / Bruno",
         text=[
             "Para resumir: primero identificamos los puntos de corte del dominio; "
             "después planteamos la condición de continuidad en cada uno; luego "
             "resolvimos el sistema de ecuaciones; y así llegamos al resultado final: "
             "a igual a siete medios, y b igual a menos tres medios.",
             "Muchas gracias por ver nuestro video. Somos Bruno Ferro, Fabrizio Osorio, "
             "Sebastián Cardoso y Eugenia Pérez. ¡Hasta la próxima!",
         ]),
]

# calcular tiempos acumulados
t = 0
for sc in SCENES:
    sc["t_start"] = t
    t += sc["dur"]
    sc["t_end"] = t
TOTAL = t


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


doc = Document()

# Márgenes
section = doc.sections[0]
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)

title = doc.add_heading("Guion — Desafío de Parámetros", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Matemática · 3° BATIN · Estación 3 · San Ignacio de Loyola School")
r.italic = True
r.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Equipo: Bruno Ferro · Fabrizio Osorio · Sebastián Cardoso · Eugenia Pérez"
)
r.font.size = Pt(11)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run(f"Duración total estimada: {fmt_time(TOTAL)} min  ({TOTAL} s)")
r2.font.size = Pt(10.5)
r2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

doc.add_paragraph()

instr = doc.add_paragraph()
instr.add_run(
    "Instrucciones de grabación: cada integrante graba su propia parte por separado "
    "(un archivo de audio .mp3/.m4a/.wav por escena), hablando con naturalidad y al ritmo "
    "indicado. Los tiempos (m:ss) marcan dónde empieza cada escena en el video ya editado; "
    "no hace falta cronometrarse de forma exacta: el script sincronizar_audios.py ajusta "
    "la duración de cada escena a la del audio grabado."
).italic = True

doc.add_paragraph()

# Tabla de reparto rápido
doc.add_heading("Reparto por escena", level=2)
table = doc.add_table(rows=1, cols=4)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
for i, h in enumerate(["#", "Escena", "Tiempo", "Integrante"]):
    hdr[i].text = h
    for run in hdr[i].paragraphs[0].runs:
        run.bold = True

for sc in SCENES:
    row = table.add_row().cells
    row[0].text = str(sc["num"])
    row[1].text = sc["title"]
    row[2].text = f'{fmt_time(sc["t_start"])}–{fmt_time(sc["t_end"])}'
    row[3].text = sc["speaker"]
    color = PERSON_COLORS[sc["speaker"]]
    for run in row[3].paragraphs[0].runs:
        run.font.color.rgb = color
        run.bold = True

doc.add_page_break()

doc.add_heading("Guion completo", level=1)

for sc in SCENES:
    color = PERSON_COLORS[sc["speaker"]]
    h = doc.add_heading(
        f'Escena {sc["num"]} — {sc["title"]}  ({fmt_time(sc["t_start"])}–{fmt_time(sc["t_end"])})',
        level=2,
    )

    sp = doc.add_paragraph()
    run = sp.add_run(f'Habla: {sc["speaker"]}')
    run.bold = True
    run.font.color.rgb = color
    run.font.size = Pt(12)

    for line in sc["text"]:
        lp = doc.add_paragraph()
        lp.paragraph_format.space_after = Pt(6)
        lrun = lp.add_run(line)
        lrun.font.color.rgb = color
        lrun.font.size = Pt(12.5)

    doc.add_paragraph()

doc.save("/home/user/Video-Mate/video_estacion3/Guion_Desafio_Parametros.docx")
print("Guion generado. Duración total:", fmt_time(TOTAL), f"({TOTAL}s)")
for sc in SCENES:
    print(sc["num"], sc["title"], fmt_time(sc["t_start"]), "-", fmt_time(sc["t_end"]), sc["speaker"])
